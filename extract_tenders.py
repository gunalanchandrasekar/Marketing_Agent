"""Phase 2a: extract text and evidence snippets from downloaded tender documents.

Run: python extract_tenders.py --input data/results_v4.json --output data/extracted_tenders.json
This stage does not infer deadlines from a search result or declare tenders open.
"""
from __future__ import annotations

import argparse
import io
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlparse

from docx import Document
from pypdf import PdfReader

DATE = r"(?:\\d{1,2}[-/. ]\\d{1,2}[-/. ]\\d{2,4}|\\d{1,2}(?:st|nd|rd|th)?\\s+[A-Za-z]{3,9}\\s+20\\d{2}|[A-Za-z]{3,9}\\s+\\d{1,2},?\\s+20\\d{2})"
FIELDS = {
    "submission_deadline": r"(?:last\\s+date(?:\\s+and\\s+time)?\\s+(?:for|of)?\\s*(?:submission|receipt|bid)|bid\\s+submission\\s+(?:end|closing)\\s+date|closing\\s+date|due\\s+date)",
    "publication_date": r"(?:date\\s+of\\s+(?:issue|publication)|published\\s+on|tender\\s+published)",
    "pre_bid_date": r"(?:pre[- ]bid\\s+(?:meeting|date)|date\\s+of\\s+pre[- ]bid)",
}
MAX_CHARS = 250_000


def read_document(path: Path) -> tuple[str, int]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        reader = PdfReader(str(path), strict=False)
        pages = len(reader.pages)
        chunks = []
        length = 0
        for page in reader.pages:
            part = page.extract_text() or ""
            chunks.append(part)
            length += len(part)
            if length >= MAX_CHARS:
                break
        return "\\n".join(chunks)[:MAX_CHARS], pages
    if suffix == ".docx":
        doc = Document(str(path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            paragraphs.extend(" | ".join(c.text for c in row.cells) for row in table.rows)
        return "\\n".join(paragraphs)[:MAX_CHARS], 0
    raise ValueError(f"Unsupported document format: {suffix}; legacy .doc needs conversion")


def evidence_candidates(text: str) -> dict[str, list[dict]]:
    """Find source quotes only, not invented structured dates."""
    result = {}
    for field, pattern in FIELDS.items():
        hits = []
        for match in re.finditer(pattern, text, re.I):
            segment = text[match.start():min(len(text), match.end()+180)]
            nearby_date = re.search(DATE, segment, re.I)
            hits.append({
                "date_candidate": nearby_date.group(0) if nearby_date else None,
                "evidence": re.sub(r"\\s+", " ", text[max(0,match.start()-35):min(len(text),match.end()+180)]).strip(),
                "offset": match.start()
            })
            if len(hits) == 10:
                break
        result[field] = hits
    return result


def extract(input_file: Path, output_file: Path) -> dict:
    source = json.loads(input_file.read_text(encoding="utf-8"))
    extracted, errors = [], []
    for item in source.get("documents", []):
        path = Path(item["path"])
        # Resolve relative paths against CWD (normal scraper behavior)
        if not path.is_file():
            errors.append({"file": str(path), "error": "File does not exist"})
            continue
        try:
            text, page_count = read_document(path)
            extracted.append({
                "document_url": item.get("url"),
                "local_path": str(path),
                "sha256": item.get("sha256"),
                "page_count": page_count,
                "text_characters": len(text),
                "extraction_status": "ok" if text.strip() else "no_embedded_text_ocr_needed",
                "date_evidence": evidence_candidates(text),
                "text": text,
            })
        except Exception as exc:
            errors.append({"file": str(path), "error": str(exc)})
    output = {
        "topic": source.get("topic"), "extracted_at": datetime.now(timezone.utc).isoformat(),
        "documents": extracted, "errors": errors
    }
    output_file.parent.mkdir(parents=True, exist_ok=True)
    output_file.write_text(json.dumps(output, ensure_ascii=False, indent=2), encoding="utf-8")
    return output


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--input", type=Path, default=Path("data/results_v4.json"))
    p.add_argument("--output", type=Path, default=Path("data/extracted_tenders.json"))
    args = p.parse_args()
    output = extract(args.input, args.output)
    for doc in output["documents"]:
        print(f'{doc["extraction_status"]}: {doc["page_count"]} pages, {doc["text_characters"]} chars: {doc["document_url"]}')
        for field, hits in doc["date_evidence"].items():
            print(f"  {field}: {len(hits)} evidence snippets")
    print(f'Extracted {len(output["documents"])} documents; errors: {len(output["errors"])}; output: {args.output}')


if __name__ == "__main__":
    main()
